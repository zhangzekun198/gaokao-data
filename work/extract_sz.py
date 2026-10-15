# -*- coding: utf-8 -*-
"""从 2026深圳中考志愿填报工具包.xlsx 抽取为深圳中考志愿填报工具的 DATA 字面量。
输出 work/sz_data.json（再由注入脚本替换 sz-zhongkao.html 中的 __DATA__ 占位）。
"""
import openpyxl, json, re, unicodedata

SRC = r"D:/BaiduNetdiskDownload/深圳中考志愿填报择校资料包/2026深圳中考志愿填报工具包.xlsx"
OUT = r"C:/Users/zhangzekun03/gaokao_data_repo/work/sz_data.json"

def norm(s):
    if s is None: return None
    s = unicodedata.normalize('NFKC', str(s)).strip()
    return s or None

def num(x):
    if x is None: return None
    s = str(x).replace(',', '').strip()
    if s in ('', '-', '--', '—', 'None'): return None
    try:
        f = float(s)
        return int(f) if f == int(f) else round(f, 1)
    except ValueError:
        return None

wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

# ---------- 公办高中 ----------
scores = []
ws = wb['公办高中']
rows = list(ws.iter_rows(values_only=True))
# 数据从 title 之后的第 3 行起（row index 2）。用首列是否为数字判断。
for r in rows:
    if len(r) < 16: continue
    idx = num(r[0])
    if idx is None: continue
    name = norm(r[2])
    if not name: continue
    ac = [num(r[6]), num(r[7])]
    d = [num(r[8]), num(r[9])]
    scores.append({
        'name': name,
        'grade': norm(r[1]),
        'district': norm(r[4]),
        'belong': norm(r[5]),
        'nature': '公办',
        'b': '第一批',
        'ac2025': ac, 'd2025': d,
        'ac2026pred': num(r[10]), 'd2026pred': num(r[12]),
        'plan2025': num(r[14]), 'plan2025_board': num(r[15]),
    })

# ---------- 民办高中 ----------
ws = wb['民办高中']
rows = list(ws.iter_rows(values_only=True))
for r in rows:
    if len(r) < 9: continue
    idx = num(r[0])
    if idx is None: continue
    name = norm(r[1])
    if not name: continue
    sc = num(r[4])
    scores.append({
        'name': name,
        'grade': None,
        'district': norm(r[3]),
        'belong': None,
        'nature': '民办',
        'b': '第一批',
        'ac2025': [sc, sc], 'd2025': [sc, sc],
        'ac2026pred': None, 'd2026pred': None,
        'plan2025': num(r[5]),
        'tuition': norm(r[8]),
        'board': norm(r[7]),
        'courses': norm(r[6]),
    })

# ---------- 指标生控制线（2024） ----------
indicator = []
ws = wb['2024指标生录取情况']
rows = list(ws.iter_rows(values_only=True))
for r in rows:
    if len(r) < 7: continue
    idx = num(r[0])
    if idx is None: continue
    name = norm(r[2])
    if not name: continue
    indicator.append({
        'name': name,
        'year': 2024,
        'ac_rec': num(r[3]), 'd_rec': num(r[4]),
        'ac_ctrl': num(r[5]), 'd_ctrl': num(r[6]),
    })

# ---------- 一类自招（原模式） ----------
zizhu = []
ws = wb['一类自招（原模式）']
rows = list(ws.iter_rows(values_only=True))
for r in rows:
    if len(r) < 14: continue
    idx = num(r[0])
    if idx is None: continue
    name = norm(r[1])
    if not name: continue
    zizhu.append({
        'name': name,
        'type': norm(r[3]),
        'dir': norm(r[4]),
        'plan': num(r[5]),
        'ac2025': [num(r[7]), num(r[8])],
        'd2025': [num(r[9]), num(r[10])],
        'ctrl': num(r[11]),
    })

# ---------- 中考等级线 ----------
lines = {}
ws = wb['中考等级线']
rows = list(ws.iter_rows(values_only=True))
for r in rows:
    if not r or len(r) < 8: continue
    key = norm(r[0])
    if not key: continue
    m = re.match(r'(\d{4})(A\+|A|B\+|B|C\+|C)线', key)
    if not m: continue
    yr, lv = m.group(1), m.group(2)
    tot = num(r[7])
    if tot is None: continue
    lines.setdefault(yr, {})[lv] = tot

# ---------- 中考时间轴（关键日程） ----------
timeline = []
ws = wb['中考时间轴']
rows = list(ws.iter_rows(values_only=True))
for r in rows:
    if len(r) < 3: continue
    content = norm(r[0]); date = norm(r[1]); event = norm(r[2])
    if not date and not event: continue
    if content and not date and not event: continue
    if date or event:
        timeline.append({'date': date, 'event': event or content})

# ---------- 组装 DATA ----------
data = {
    'meta': {
        'city': '深圳', 'year': 2025, 'total': 610,
        'src': '2026深圳中考志愿填报工具包.xlsx（公办/民办高中、指标生、自主招生）',
        'note': '录取分为 2025 年最低录取分（AC类=深户及不限户籍 / D类=非深户），含住宿/走读两条线；智能推荐取住宿线定位。2026 预测线（26指控测）已并入。精确位次请以官方“一分一段表”为准。',
        'year2026': '2026 预测线按 2025 录取分 + 专家降分测算并入',
    },
    'policy': {
        'batches': [
            {'name': '自主招生批', 'vol': '1 个志愿', 'note': '普通高中一类 / 二类自招，按特长 / 专项考核择优，在中考成绩基础上合成总分排序'},
            {'name': '指标生批', 'vol': '1 个志愿', 'note': '原“名额分配”，户籍生专属降分渠道；达到该校指标生控制线即可在同校竞争中降分录取'},
            {'name': '第一批（普通高中）', 'vol': '12 普高 + 4 中职志愿', 'note': '按分录取主力批次，AC 类（深户及不限户籍）/ D 类（非深户）分别划线投档'},
            {'name': '第二批（中职）', 'vol': '18 个中职志愿', 'note': '本市中等职业学校，可选不超过 4 所学校勾选“接受走读调剂”'},
            {'name': '第三批（部省属中职）', 'vol': '6 个志愿', 'note': '部、省属中等职业学校在深招生计划'},
        ],
        'lines': lines,
        'rule': '深圳中考实行“分数优先、依照志愿顺序”投档：在同一批次内，按考生总分从高到低依次检索其志愿，分数达标即投档，遵循志愿顺序。被前一批次录取即不再参加后续批次。',
        'hukou': 'ACD 类：AC = 深户（及不限户籍的计划），D = 非深户。AC 类、D 类分别划定最低录取控制线，D 类分数线通常高于 AC 类。非深户报考公办高中须满足深圳市随迁子女升学条件（父母持有效居住证、合法稳定职业与住所、社保缴纳年限等）。',
        'minge': '指标生（原名额分配）面向具有深圳户籍、在同一初中学校连续就读三年且拥有正式学籍的应届初三毕业生。考生达到招生学校指标生控制线，即可以低于正取线一定的分数在同校（同初中）竞争中优先录取。指标生批设 1 个志愿。',
        'principle': '志愿之间必须拉开梯度：每个志愿的预估录取分应逐级下降，避免全部扎在同一分数段导致滑档。可参考“冲—稳—保”结构，结合各校近年最低分与 2026 预测线排序。',
        'total': 610,
    },
    'scores': scores,
    'indicator': indicator,
    'zizhu': zizhu,
    'timeline': timeline,
}

with open(OUT, 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=1)

print('=== 抽取统计 ===')
print('公办+民办高中(scores):', len(scores), '公办:', sum(1 for s in scores if s['nature']=='公办'), '民办:', sum(1 for s in scores if s['nature']=='民办'))
print('指标生(indicator):', len(indicator))
print('自主招生(zizhu):', len(zizhu))
print('等级线(lines)年份:', list(lines.keys()))
print('时间轴(timeline):', len(timeline))
print('输出:', OUT)
# 抽样校验
print('--- 公办样例 ---', json.dumps(scores[0], ensure_ascii=False))
print('--- 民办样例 ---', json.dumps([s for s in scores if s['nature']=='民办'][0], ensure_ascii=False))
print('--- 指标生样例 ---', json.dumps(indicator[0], ensure_ascii=False))
print('--- 自招样例 ---', json.dumps(zizhu[0], ensure_ascii=False))
